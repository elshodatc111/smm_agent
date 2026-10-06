"""Hisobotlar: ekranda (data), PDF (ReportLab) va Excel (openpyxl); agentlik brendi bilan."""
from __future__ import annotations

import io
from datetime import date, datetime, time, timedelta

from django.conf import settings
from django.core.files.base import ContentFile
from django.utils import timezone

from apps.core.settings_store import get_setting
from apps.core.utils import fmt_num, tz

FONT_DIR = settings.BASE_DIR / "static" / "fonts"


def norm(t) -> str:
    return str(t or "").replace("ʻ", "'").replace("ʼ", "'").replace("’", "'").replace("‘", "'")


def collect_data(report) -> dict:
    from apps.metrics import analytics as an
    from apps.metrics.models import CompetitorMetrics
    b = report.business
    since = datetime.combine(report.period_start, time.min, tzinfo=tz())
    until = datetime.combine(report.period_end + timedelta(days=1), time.min, tzinfo=tz())
    days = max(1, (report.period_end - report.period_start).days + 1)
    cur = an.period_kpis(b, since, until)
    prev = an.period_kpis(b, since - (until - since), since)
    delta = {k: ((v - prev[k]) / prev[k] * 100) if isinstance(v, (int, float)) and isinstance(prev.get(k), (int, float)) and prev[k] else None for k, v in cur.items()}
    tops = [{"text": (r.post.text or "")[:90].replace("\n", " "), "type": r.post.media_type, "views": r.views, "er": r.er,
             "platform": r.post.platform, "url": r.post.url} for r in an.top_posts(b, 5, days=days)]
    worst = [{"text": (r.post.text or "")[:90].replace("\n", " "), "views": r.views, "er": r.er} for r in an.top_posts(b, 3, worst=True, days=days)]
    series = an.business_follower_series(b, days)
    comps = []
    for c in b.competitors.filter(active=True):
        ms = list(CompetitorMetrics.objects.filter(competitor=c).order_by("-taken_at")[:2])
        if ms:
            first = CompetitorMetrics.objects.filter(competitor=c, taken_at__gte=since).order_by("taken_at").first() or ms[-1]
            comps.append({"name": c.name or c.handle, "platform": c.platform, "followers": ms[0].followers,
                          "growth": (ms[0].followers or 0) - (first.followers or 0)})
    from apps.ads.models import AdCampaign
    ads = []
    if report.kind in ("ads", "monthly"):
        for c in AdCampaign.objects.filter(business=b, status__in=["running", "done"]):
            from apps.ads.services import campaign_variance
            ads.append({"name": c.name, "variance": campaign_variance(c)})
    out = {"kpi": cur, "prev": prev, "delta": delta, "top": tops, "worst": worst, "by_type": an.by_type(b, days),
           "series": series, "competitors": comps, "ads": ads, "frequency": an.post_frequency(b, days),
           "best_times": {k: v[:3] for k, v in an.best_times(b).items()}}
    out.update(extra_data(report, since, until))
    return out


def extra_data(report, since, until) -> dict:
    """Mijoz hisoboti uchun qo'shimcha: platformalar kesimi, kunlik statistika, nashrlar, stories."""
    from apps.metrics import analytics as an
    from apps.metrics.daily import rollup_business
    from apps.metrics.models import DailyStat, Post, Story
    b = report.business
    # platformalar kesimi
    per = []
    for plat, name in (("telegram", "Telegram"), ("instagram", "Instagram"), ("youtube", "YouTube")):
        ids = list(b.accounts.filter(platform=plat, archived_at__isnull=True).values_list("pk", flat=True))
        if not ids:
            continue
        k = an.period_kpis(b, since, until, ids)
        pk = an.period_kpis(b, since - (until - since), since, ids)
        per.append({"platform": plat, "name": name, "kpi": k, "prev_followers": pk.get("followers")})
    # kunlik
    d0, d1 = report.period_start, report.period_end
    rows = list(DailyStat.objects.filter(business=b, account__isnull=True, day__gte=d0, day__lte=d1).order_by("day"))
    if not rows:
        try:
            cur = d0
            while cur <= min(d1, timezone.localdate()) and (cur - d0).days < 92:
                rollup_business(b, cur)
                cur += timedelta(days=1)
            rows = list(DailyStat.objects.filter(business=b, account__isnull=True, day__gte=d0, day__lte=d1).order_by("day"))
        except Exception:  # noqa: BLE001
            rows = []
    daily = [{"day": r.day.isoformat(), "followers": r.followers, "delta": r.followers_delta, "views": r.views, "posts": r.posts, "stories": r.stories,
              "clicks": r.link_clicks, "story_views": r.story_views} for r in rows]
    clicks = sum(r["clicks"] for r in daily)
    daily_total = {"days": len(daily), "followers_delta": sum(r["delta"] or 0 for r in daily), "views": sum(r["views"] for r in daily),
                   "posts": sum(r["posts"] for r in daily), "stories": sum(r["stories"] for r in daily), "clicks": clicks}
    # nashrlar ro'yxati
    pubs = []
    for p in Post.objects.filter(business=b, published_at__gte=since, published_at__lt=until).select_related("account").order_by("-published_at")[:40]:
        pubs.append({"when": p.published_at.astimezone(tz()).strftime("%d.%m %H:%M"), "platform": p.platform, "type": p.media_type, "text": (p.text or "")[:80].replace("\n", " "), "url": p.url})
    st = Story.objects.filter(business=b, published_at__gte=since, published_at__lt=until)
    sv = sum(((h.views or 0) if (h := s.history.order_by("-is_final", "-taken_at").first()) else 0) for s in st)
    return {"per_platform": per, "daily": daily, "daily_total": daily_total, "clicks": clicks, "publications": pubs, "stories": {"count": st.count(), "views": sv}}


def ai_summary(report, data: dict, router=None) -> dict:
    """AI xulosa + keyingi oy tavsiyasi. AI yo'q bo'lsa qoidaga asoslangan xulosa."""
    k, d = data["kpi"], data["delta"]
    fallback = (f"Davr: {report.period_start:%d.%m.%Y} – {report.period_end:%d.%m.%Y}. Obunachilar: {fmt_num(k['followers'])} "
                f"({k['followers_delta']:+d}). Postlar: {k['posts']}, ko'rishlar: {fmt_num(k['views'])}, "
                f"ER: {(k['er'] or 0):.1%}.")
    recs = []
    if data["by_type"]:
        recs.append(f"Eng samarali format: {data['by_type'][0]['type']} — bunday kontentni ko'paytiring.")
    if data["worst"]:
        recs.append("Eng past natija bergan postlar mavzusi/hookini qayta ko'rib chiqing.")
    try:
        import json
        from apps.ai_router.router import generate
        from apps.ai_router.types import TaskType
        prompt = ("Mijoz uchun hisobot xulosasi: 4-6 gap, sodda o'zbekcha, raqamlarga tayanib. Oxirida 'Nimani o'zgartirish kerak' va 3 ta tavsiya. "
                  'JSON: {"summary":"","recommendations":[]}\nRaqamlar: ' + json.dumps({"kpi": k, "delta": d, "top": data["top"][:3], "types": data["by_type"][:3]}, ensure_ascii=False, default=str))
        res, _ = generate(TaskType.LONG_UZ, prompt, business=report.business, kind="report", schema={"json": True}, router=router)
        if res.data and res.data.get("summary"):
            return {"summary": res.data["summary"], "recommendations": res.data.get("recommendations") or recs, "ai": True}
    except Exception:  # noqa: BLE001
        pass
    return {"summary": fallback, "recommendations": recs, "ai": False}


def _brand_color():
    return get_setting("reports.brand_color", "#E11D2E")


def build_pdf(report, data: dict, summary: dict) -> bytes:
    from .design import build_pdf as _bp
    return _bp(report, data, summary)


def build_xlsx(report, data: dict, summary: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    hdr = PatternFill("solid", fgColor=_brand_color().lstrip("#"))
    white = Font(color="FFFFFF", bold=True)

    def sheet(name, header, rows, first=False):
        ws = wb.active if first else wb.create_sheet()
        ws.title = name
        ws.append(header)
        for c in ws[1]:
            c.fill, c.font = hdr, white
        for r in rows:
            ws.append(r)
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = min(60, max(12, max(len(str(c.value or "")) for c in col) + 2))
        return ws
    k = data["kpi"]
    sheet("Umumiy", ["Ko'rsatkich", "Qiymat"], [["Obunachilar", k["followers"]], ["O'sish", k["followers_delta"]], ["Postlar", k["posts"]],
                                              ["Ko'rishlar", k["views"]], ["Qamrov", k["reach"]], ["Layk", k["likes"]], ["Izoh", k["comments"]],
                                              ["Ulashish", k["shares"]], ["Saqlash", k["saves"]], ["ER", k["er"]], ["Xulosa", summary["summary"]]], first=True)
    sheet("Eng yaxshi postlar", ["Matn", "Tur", "Tarmoq", "Ko'rish", "ER", "Havola"], [[p["text"], p["type"], p["platform"], p["views"], p["er"], p["url"]] for p in data["top"]])
    sheet("Turlar", ["Tur", "Soni", "O'rtacha ko'rish", "O'rtacha ER"], [[t["type"], t["count"], t["avg_views"], t["avg_er"]] for t in data["by_type"]])
    sheet("Obunachilar", ["Sana", "Obunachi"], [[s["date"], s["followers"]] for s in data["series"]])
    sheet("Raqobatchilar", ["Nomi", "Tarmoq", "Obunachi", "O'sish"], [[c["name"], c["platform"], c["followers"], c["growth"]] for c in data["competitors"]])
    sheet("Kunlik", ["Sana", "Obunachi", "O'sish", "Postlar", "Stories", "Ko'rish", "Story ko'rish", "Havola bosish"],
          [[x["day"], x["followers"], x["delta"], x["posts"], x["stories"], x["views"], x["story_views"], x["clicks"]] for x in data.get("daily", [])])
    sheet("Platformalar", ["Platforma", "Obunachi", "O'sish", "Postlar", "Ko'rish", "Layk", "ER"],
          [[x["name"], x["kpi"]["followers"], x["kpi"]["followers_delta"], x["kpi"]["posts"], x["kpi"]["views"], x["kpi"]["likes"], x["kpi"]["er"]] for x in data.get("per_platform", [])])
    sheet("Tavsiyalar", ["Tavsiya"], [[r] for r in summary["recommendations"]])
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def generate_report(report, router=None) -> None:
    try:
        data = collect_data(report)
        summary = ai_summary(report, data, router)
        report.data = {**{k: v for k, v in data.items() if k in ("kpi", "delta", "top", "frequency", "per_platform", "daily", "daily_total", "clicks", "publications", "stories", "by_type", "competitors")}, "summary": summary}
        report.summary_text = summary["summary"] + "\n\n" + "\n".join("• " + r for r in summary["recommendations"])
        stamp = f"{report.business_id}_{report.kind}_{report.period_end:%Y%m%d}"
        report.pdf.save(f"{stamp}.pdf", ContentFile(build_pdf(report, data, summary)), save=False)
        report.xlsx.save(f"{stamp}.xlsx", ContentFile(build_xlsx(report, data, summary)), save=False)
        report.status, report.error = "ready", ""
    except Exception as e:  # noqa: BLE001
        report.status, report.error = "failed", f"{type(e).__name__}: {e}"[:300]
        raise
    finally:
        report.save()


def default_period(kind: str, today: date | None = None) -> tuple[date, date]:
    today = today or timezone.now().astimezone(tz()).date()
    if kind == "weekly":
        end = today - timedelta(days=today.weekday() + 1)
        return end - timedelta(days=6), end
    if kind == "monthly":
        end = today.replace(day=1) - timedelta(days=1)
        return end.replace(day=1), end
    return today - timedelta(days=29), today
