"""Onlayn kanal bo'limi: loyihalar (bot + yopiq guruh), bot konstruktori, leadlar, obunachilar, to'lovlar, tariflar, AI mentor. Biznesga bog'liq emas."""
from __future__ import annotations

import io
import json
import re
from datetime import timedelta

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.paginator import Paginator
from django.db.models import Count, Q
from django.http import FileResponse, Http404, HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone
from django.views.decorators.http import require_POST

from apps.accounts.decorators import require_permission
from apps.accounts.services import perms
from apps.core.audit import audit
from apps.core.utils import fmt_dt
from apps.online import ai as mentor
from apps.online import defaults as D
from apps.online import services as S
from apps.online.bot import send_screen
from apps.online.models import (LEAD_STATUS, KnowledgeDoc, Lead, LeadNote, OnlineMedia, OnlineProject, Payment, ProjectAdmin, ProjectGroup, Subscription,
                                Tariff)
from apps.online.tgapi import TgError, get_api

BADGE = {"new": "", "phone": "blue", "interested": "violet", "waiting": "warn", "sold": "ok", "lost": "bad",
         "pending": "warn", "approved": "ok", "rejected": "bad", "active": "ok", "expired": "", "cancelled": "bad",
         "draft": "warn", "paused": ""}
IMG_EXT = ("jpg", "jpeg", "png", "webp")
VID_EXT = ("mp4", "mov", "m4v")


def _can(request, code):
    return perms.can(request.user, code, None)


def _proj(pk) -> OnlineProject:
    return get_object_or_404(OnlineProject, pk=pk)


def _tabs(request, p):
    allt = [("online_project", "Umumiy", "online.view"), ("online_bot", "Bot", "online.view"), ("online_leads", "Leadlar", "online.leads_view"),
            ("online_subs", "Obunachilar", "online.view"), ("online_payments", "To'lovlar", "online.view"), ("online_tariffs", "Tariflar", "online.view"),
            ("online_ai", "AI mentor", "online.view"), ("online_groups", "Guruhlar va adminlar", "online.view"), ("online_events", "Tarix", "online.view"),
            ("online_settings", "Sozlamalar", "online.configure")]
    return [(u, l) for u, l, c in allt if _can(request, c)]


def _ctx(request, p, tab, **kw):
    return {"pbadge": BADGE.get(p.status, ""), "p": p, "tabs": _tabs(request, p), "tab": tab, "badge": BADGE, "title": f"{p.name} · Onlayn kanal", **kw}


def _api(p):
    return get_api(p)


def _back(request, p, name="online_subs"):
    ref = request.META.get("HTTP_REFERER")
    if ref and ref.startswith(("/", request.build_absolute_uri("/"))):
        return redirect(ref)
    return redirect(name, p.pk)


def _by(request):
    return f"{request.user.get_username()} (panel)"


# ───────────────────────── ro'yxat va yaratish ─────────────────────────
@login_required
@require_permission("online.view")
def online_home(request):
    rows = []
    for p in OnlineProject.objects.all():
        rows.append({"p": p, "s": S.stats(p, 7), "all": S.stats(p)})
    return render(request, "online_home.html", {"title": "Onlayn kanal", "rows": rows, "can_create": _can(request, "online.configure"), "badge": BADGE})


@login_required
@require_permission("online.configure")
def online_new(request):
    if request.method == "POST":
        P = request.POST
        name = P.get("name", "").strip()[:120]
        if not name:
            messages.error(request, "Loyiha nomini kiriting.")
            return render(request, "online_new.html", {"title": "Yangi loyiha", "P": P})
        p = OnlineProject(name=name, brief=P.get("brief", "")[:3000], created_by=request.user)
        p.set_token(P.get("token", ""))
        p.flow = D.default_flow()
        p.bot_title = name[:64]
        p.save()
        audit("online.project_created", user=request.user, target=p.name)
        if P.get("token", "").strip():
            try:
                me = _api(p).get_me()
                p.bot_id, p.bot_username = me["id"], me.get("username", "")
                p.save(update_fields=["bot_id", "bot_username"])
            except TgError as e:
                messages.warning(request, f"Token tekshirilmadi: {e.description}. Sozlamalarda to'g'rilang.")
        Tariff.objects.create(project=p, name="1 oylik", days=30, price=0, description="", sort=1)
        if P.get("ai") and (P.get("brief") or "").strip():
            try:
                r = mentor.generate_bot(p, P["brief"], user=request.user)
                messages.success(request, "AI bot matnlarini tayyorladi (qoralamaga yozildi): " + ", ".join(r["done"] or ["javob bo'sh"]))
            except Exception as e:  # noqa: BLE001
                messages.warning(request, f"AI matnlarni tayyorlay olmadi: {getattr(e, 'user_message_uz', None) or e}")
        messages.success(request, "Loyiha yaratildi. Endi tarif, karta, guruh va adminni kiriting, so'ng «Botni sozlash» tugmasini bosing.")
        return redirect("online_project", p.pk)
    return render(request, "online_new.html", {"title": "Yangi loyiha", "P": {}})


# ───────────────────────── loyiha paneli ─────────────────────────
@login_required
@require_permission("online.view")
def online_project(request, pk):
    p = _proj(pk)
    periods = [("Bugun", S.stats(p, 1)), ("7 kun", S.stats(p, 7)), ("30 kun", S.stats(p, 30)), ("Hammasi", S.stats(p))]
    series = S.daily_series(p, 30)
    mx = max([r["leads"] for r in series] + [r["sales"] for r in series] + [1])
    for r in series:
        r["lh"], r["sh"] = int(r["leads"] * 100 / mx), int(r["sales"] * 100 / mx)
    pending = p.payments.filter(status="pending").select_related("lead", "tariff")[:10]
    todo = []
    if not p.token_enc:
        todo.append("Bot tokenini kiriting (Sozlamalar).")
    if not p.admins.exists():
        todo.append("Mas'ul adminni biriktiring (Guruhlar va adminlar).")
    if not p.groups.exists():
        todo.append("Yopiq guruhni ulang.")
    if not p.card_number:
        todo.append("To'lov kartasini kiriting (Sozlamalar).")
    if not p.tariffs.filter(is_active=True, price__gt=0).exists():
        todo.append("Kamida bitta narxli tarif qo'shing.")
    if not p.leads_chat_id:
        todo.append("Leadlar guruhini ulang.")
    return render(request, "online_project.html", _ctx(request, p, "online_project", periods=periods, series=series, pending=pending, todo=todo,
                                                       can_publish=_can(request, "online.publish"), can_pay=_can(request, "online.payments")))


@login_required
@require_POST
@require_permission("online.publish")
def online_publish(request, pk):
    p = _proj(pk)
    checks = S.publish_bot(p)
    bad = [c for c in checks if not c["ok"]]
    audit("online.bot_published", user=request.user, target=p.name)
    messages.success(request, f"Bot sozlandi: {len(checks) - len(bad)}/{len(checks)} tekshiruv o'tdi." + (" Quyida muammolar ko'rsatilgan." if bad else ""))
    return redirect("online_project", p.pk)


@login_required
@require_POST
@require_permission("online.publish")
def online_status(request, pk):
    p = _proj(pk)
    if p.status == "active":
        p.status = "paused"
    elif p.token_enc:
        p.status = "active"
    p.save(update_fields=["status"])
    S.event(p, "status", f"Bot holati: {p.get_status_display()} ({request.user.get_username()})")
    return redirect("online_project", p.pk)


# ───────────────────────── bot konstruktori ─────────────────────────
@login_required
@require_permission("online.view")
def online_bot(request, pk):
    p = _proj(pk)
    flow = D.get_flow(p, draft=True)
    media = [{"id": m.pk, "url": f"/onlayn/{p.pk}/media/{m.pk}/", "kind": m.kind} for m in p.media.all()]
    t0 = p.tariffs.filter(is_active=True).first()
    sample = S.user_vals(p, Lead(first_name="Ali", tg_id=0, project=p), t0)
    sample.update(name="Ali", end="31 dekabr 2026", left=30, reason="Chek aniq emas", name_="")
    cmds = "\n".join(f"{c['command']} - {c['description']}" for c in (p.commands or S.DEFAULT_COMMANDS))
    return render(request, "online_bot.html", _ctx(request, p, "online_bot", flow=flow, media=media, sample={k: str(v) for k, v in sample.items()}, actions=D.ACTIONS,
                                                   has_draft=bool(p.flow_draft), cmds=cmds, versions=p.versions.all()[:8], avatar=p.avatar.url if p.avatar else "",
                                                   can_edit=_can(request, "online.edit"), can_publish=_can(request, "online.publish"),
                                                   tariffs=[{"id": t.pk, "name": t.name, "price": D.money(t.price), "days": t.days, "description": t.description} for t in p.tariffs.filter(is_active=True)],
                                                   admins=p.admins.all(), placeholders=D.PLACEHOLDERS, text_labels=TEXT_LABELS,
                                                   profile={"title": p.bot_title or p.name, "about": p.about, "description": p.description, "avatar": p.avatar.url if p.avatar else ""}))


TEXT_LABELS = {
    "tariffs_title": "Tariflar ro'yxati sarlavhasi", "no_tariffs": "Tarif yo'q bo'lganda", "pay_info": "To'lov rekvizitlari xabari",
    "pay_no_card": "Karta kiritilmagan bo'lsa", "pay_wait": "Chek yuborilgandan keyin", "need_photo": "Chek o'rniga matn yuborilsa",
    "approved": "To'lov tasdiqlanganda", "rejected": "To'lov rad etilganda", "expiring": "Obuna tugashidan oldingi eslatma",
    "expired": "Obuna tugaganda", "no_sub": "Obuna yo'q bo'lganda", "status_active": "Obuna holati (faol)", "join_links": "Guruh havolalari sarlavhasi",
    "no_groups": "Guruh ulanmagan bo'lsa", "join_denied": "Obunasiz guruhga kirishga urinsa", "phone_thanks": "Telefon qabul qilinganda",
    "phone_bad": "Telefon noto'g'ri yuborilsa", "ai_intro": "AI mentor ulanganda", "ai_limit": "AI kunlik limit tugaganda",
    "ai_off": "AI o'chirilgan bo'lsa", "ai_error": "AI xatosi", "blocked": "Qora ro'yxatdagi foydalanuvchiga", "stat_denied": "Statistika (admin emas)",
    "unknown": "Tushunarsiz xabar",
}


@login_required
@require_POST
@require_permission("online.edit")
def online_bot_save(request, pk):
    p = _proj(pk)
    try:
        data = json.loads(request.body.decode() or "{}")
        flow, warns = D.clean_flow(data, set(p.media.values_list("pk", flat=True)))
    except (ValueError, json.JSONDecodeError) as e:
        return JsonResponse({"ok": False, "error": str(e)}, status=400)
    p.flow_draft = flow
    p.save(update_fields=["flow_draft"])
    return JsonResponse({"ok": True, "warns": warns, "saved": fmt_dt(timezone.now())})


@login_required
@require_POST
@require_permission("online.publish")
def online_bot_flow_publish(request, pk):
    p = _proj(pk)
    if not p.flow_draft:
        messages.info(request, "Qoralamada o'zgarish yo'q.")
    else:
        S.publish_flow(p, request.user)
        audit("online.flow_published", user=request.user, target=p.name)
        messages.success(request, "Bot oqimi joriy qilindi: foydalanuvchilar endi yangi matn va tugmalarni ko'radi.")
    return redirect("online_bot", p.pk)


@login_required
@require_POST
@require_permission("online.edit")
def online_bot_discard(request, pk):
    p = _proj(pk)
    p.flow_draft = None
    p.save(update_fields=["flow_draft"])
    messages.success(request, "Qoralama bekor qilindi, joriy oqim ko'rsatilmoqda.")
    return redirect("online_bot", p.pk)


@login_required
@require_POST
@require_permission("online.publish")
def online_bot_rollback(request, pk):
    p = _proj(pk)
    v = get_object_or_404(p.versions, pk=request.POST.get("version"))
    p.flow_draft = v.data
    p.save(update_fields=["flow_draft"])
    messages.success(request, "Eski versiya qoralamaga yuklandi. Ko'rib chiqib «Joriy qilish» ni bosing.")
    return redirect("online_bot", p.pk)


@login_required
@require_POST
@require_permission("online.publish")
def online_bot_test(request, pk):
    """Qoralama ekranlarini mas'ul adminning Telegramiga yuboradi (haqiqiy ko'rinishni tekshirish)."""
    p = _proj(pk)
    adm = p.admins.filter(pk=request.POST.get("admin")).first() or p.admins.first()
    api = _api(p)
    if not (adm and api):
        messages.error(request, "Sinov uchun bot tokeni va mas'ul admin (Telegram ID) kerak.")
        return redirect("online_bot", p.pk)
    flow = D.get_flow(p, draft=True)
    lead = Lead(project=p, tg_id=adm.tg_id, first_name=adm.name or "Admin")
    try:
        api.send(adm.tg_id, "🧪 <b>Sinov:</b> qoralama ekranlari yuborilmoqda. Adminlarga mo'ljallangan tugmalar ham ko'rinadi.")
        for k in flow["screens"]:
            send_screen(api, p, adm.tg_id, k, lead, True, flow=flow)
        messages.success(request, f"{len(flow['screens'])} ta ekran {adm.name or adm.tg_id} ga yuborildi. Eslatma: adminning botga avval /start bosgan bo'lishi kerak.")
    except TgError as e:
        messages.error(request, f"Yuborib bo'lmadi: {e.description}. Admin botga /start bosganmi?")
    return redirect("online_bot", p.pk)


@login_required
@require_POST
@require_permission("online.edit")
def online_media_upload(request, pk):
    p = _proj(pk)
    f = request.FILES.get("file")
    if not f:
        return JsonResponse({"ok": False, "error": "Fayl tanlanmadi."}, status=400)
    ext = f.name.rsplit(".", 1)[-1].lower() if "." in f.name else ""
    if ext not in IMG_EXT + VID_EXT:
        return JsonResponse({"ok": False, "error": "Faqat rasm (jpg, png, webp) yoki video (mp4, mov) yuklang."}, status=400)
    if f.size > (50 if ext in VID_EXT else 10) * 1024 * 1024:
        return JsonResponse({"ok": False, "error": "Fayl katta: rasm 10 MB, video 50 MB gacha (Telegram bot cheklovi)."}, status=400)
    m = OnlineMedia.objects.create(project=p, file=f, kind="video" if ext in VID_EXT else "photo")
    return JsonResponse({"ok": True, "id": m.pk, "url": f"/onlayn/{p.pk}/media/{m.pk}/", "kind": m.kind})


@login_required
@require_permission("online.view")
def online_media_file(request, pk, mid):
    m = get_object_or_404(OnlineMedia, pk=mid, project_id=pk)
    return FileResponse(m.file.open("rb"), content_type="video/mp4" if m.kind == "video" else "image/jpeg")


@login_required
@require_POST
@require_permission("online.edit")
def online_profile(request, pk):
    p = _proj(pk)
    P = request.POST
    p.bot_title = P.get("bot_title", "").strip()[:64]
    p.about = P.get("about", "").strip()[:120]
    p.description = P.get("description", "").strip()[:512]
    p.ask_phone = bool(P.get("ask_phone"))
    cmds = []
    for line in P.get("commands", "").splitlines():
        m = re.match(r"^/?([A-Za-z0-9_]{1,32})\s*[-–:]\s*(.{1,256})$", line.strip())
        if m:
            cmds.append({"command": m.group(1).lower(), "description": m.group(2)[:256]})
    p.commands = cmds[:20]
    if request.FILES.get("avatar"):
        f = request.FILES["avatar"]
        if f.name.rsplit(".", 1)[-1].lower() in IMG_EXT and f.size < 8 * 1024 * 1024:
            p.avatar = f
        else:
            messages.error(request, "Bot rasmi uchun jpg/png (8 MB gacha) yuklang.")
    p.save()
    messages.success(request, "Bot profili saqlandi. Telegramga yuborish uchun «Botni sozlash va ishga tushirish» tugmasini bosing.")
    return redirect("online_bot", p.pk)


@login_required
@require_POST
@require_permission("online.edit")
def online_ai_generate(request, pk):
    p = _proj(pk)
    brief = request.POST.get("brief", "").strip()
    if len(brief) < 20:
        messages.error(request, "Loyiha haqida kamida 2-3 gap yozing (nima o'rgatasiz, kimlar uchun, ohang).")
        return redirect("online_bot", p.pk)
    try:
        r = mentor.generate_bot(p, brief, user=request.user)
        messages.success(request, "AI tayyorladi: " + (", ".join(r["done"]) or "javob bo'sh keldi") + ". Hammasi QORALAMAga yozildi — ko'rib chiqib joriy qiling.")
    except Exception as e:  # noqa: BLE001
        messages.error(request, getattr(e, "user_message_uz", None) or f"AI javob bermadi: {e}")
    return redirect("online_bot", p.pk)


# ───────────────────────── leadlar ─────────────────────────
def _lead_qs(request, p):
    g = request.GET
    qs = p.leads.select_related("assigned_to")
    if g.get("status"):
        qs = qs.filter(status=g["status"])
    if g.get("phone") == "1":
        qs = qs.exclude(phone="")
    if g.get("phone") == "0":
        qs = qs.filter(phone="")
    if g.get("src"):
        qs = qs.filter(source=g["src"])
    if g.get("mine") == "1":
        qs = qs.filter(assigned_to=request.user)
    q = (g.get("q") or "").strip()
    if q:
        qs = qs.filter(Q(first_name__icontains=q) | Q(last_name__icontains=q) | Q(username__icontains=q.lstrip("@")) | Q(phone__icontains=q) | Q(tg_id__icontains=q))
    return qs


@login_required
@require_permission("online.leads_view")
def online_leads(request, pk):
    p = _proj(pk)
    qs = _lead_qs(request, p)
    if request.GET.get("export") == "1":
        if not _can(request, "online.export"):
            return redirect("online_leads", p.pk)
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "Leadlar"
        ws.append(["Ism", "Username", "Telegram ID", "Telefon", "Holat", "Manba", "Til", "Premium", "Birinchi kirish", "Oxirgi faollik", "Biriktirilgan"])
        for l in qs:
            ws.append([l.full_name, l.username, l.tg_id, l.phone, l.get_status_display(), l.source, l.language, "ha" if l.is_premium else "", fmt_dt(l.first_seen),
                       fmt_dt(l.last_seen), l.assigned_to.get_username() if l.assigned_to else ""])
        buf = io.BytesIO()
        wb.save(buf)
        r = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        r["Content-Disposition"] = f'attachment; filename="leadlar_{p.pk}.xlsx"'
        return r
    page = Paginator(qs, 50).get_page(request.GET.get("page"))
    srcs = list(p.leads.exclude(source="").values_list("source", flat=True).distinct()[:30])
    counts = dict(p.leads.values_list("status").annotate(n=Count("id")))
    return render(request, "online_leads.html", _ctx(request, p, "online_leads", page=page, statuses=LEAD_STATUS, srcs=srcs, counts=counts, G=request.GET,
                                                     can_export=_can(request, "online.export"), total=p.leads.count()))


@login_required
@require_permission("online.leads_view")
def online_lead(request, pk, lid):
    p = _proj(pk)
    lead = get_object_or_404(Lead, pk=lid, project=p)
    if request.method == "POST":
        if not _can(request, "online.leads_edit"):
            return redirect("online_lead", p.pk, lead.pk)
        a = request.POST.get("action")
        if a == "status" and request.POST.get("status") in dict(LEAD_STATUS):
            S.set_lead_status(p, _api(p), lead, request.POST["status"])
        elif a == "note" and request.POST.get("text", "").strip():
            LeadNote.objects.create(lead=lead, user=request.user, text=request.POST["text"].strip()[:2000])
        elif a == "assign":
            from apps.accounts.models import User
            lead.assigned_to = User.objects.filter(pk=request.POST.get("user") or 0, is_active=True).first()
            lead.save(update_fields=["assigned_to"])
        elif a == "block":
            lead.blocked = not lead.blocked
            lead.save(update_fields=["blocked"])
            S.event(p, "block", f"{lead.full_name}: {'qora ro‘yxatga qo‘shildi' if lead.blocked else 'qora ro‘yxatdan chiqarildi'}", lead)
        elif a == "repost":
            lead.posted_at = None
            lead.save(update_fields=["posted_at"])
            S.post_lead(p, _api(p), lead, activity=True)
            messages.success(request, "Lead leadlar guruhiga qayta yuborildi.")
        return redirect("online_lead", p.pk, lead.pk)
    from apps.accounts.models import User
    return render(request, "online_lead.html", _ctx(request, p, "online_leads", lead=lead, statuses=LEAD_STATUS, users=User.objects.filter(is_active=True).order_by("username"),
                                                    pays=lead.payments.select_related("tariff"), subs=lead.subs.select_related("tariff")[:5], events=lead.events.all()[:30],
                                                    can_edit=_can(request, "online.leads_edit"), sub=S.active_sub(lead)))


# ───────────────────────── obunachilar ─────────────────────────
def _find_lead(p, q):
    q = (q or "").strip()
    if not q:
        return None
    qs = p.leads.all()
    if q.lstrip("@").isdigit() and len(q) < 13 and not q.startswith("+"):
        l = qs.filter(tg_id=int(q)).first()
        if l:
            return l
    ph = S.norm_phone(q) if re.sub(r"\D", "", q) and not q.startswith("@") else ""
    return qs.filter(Q(username__iexact=q.lstrip("@")) | (Q(phone=ph) if ph else Q(pk__in=[]))).first()


@login_required
@require_permission("online.view")
def online_subs(request, pk):
    p = _proj(pk)
    f = request.GET.get("f", "active")
    now = timezone.now()
    qs = p.subs.select_related("lead", "tariff")
    if f == "active":
        qs = qs.filter(status="active", ends_at__gt=now).order_by("ends_at")
    elif f == "soon":
        qs = qs.filter(status="active", ends_at__gt=now, ends_at__lte=now + timedelta(days=7)).order_by("ends_at")
    elif f == "ended":
        qs = qs.exclude(status="active")
    page = Paginator(qs, 50).get_page(request.GET.get("page"))
    exempt = p.leads.filter(exempt=True)
    return render(request, "online_subs.html", _ctx(request, p, "online_subs", page=page, f=f, exempt=exempt, can_members=_can(request, "online.members"),
                                                    n_active=p.subs.filter(status="active", ends_at__gt=now).count()))


@login_required
@require_POST
@require_permission("online.members")
def online_sub_action(request, pk):
    p = _proj(pk)
    api = _api(p)
    P = request.POST
    a = P.get("action")
    lead = p.leads.filter(pk=P.get("lead") or 0).first() or _find_lead(p, P.get("q"))
    if not lead:
        messages.error(request, "Lead topilmadi. Telegram ID, @username yoki telefon raqamini kiriting (u botga /start bosgan bo'lishi kerak).")
        return _back(request, p)
    try:
        days = max(1, min(3650, int(P.get("days") or 30)))
    except ValueError:
        days = 30
    if a in ("grant", "extend"):
        sub = S.grant_subscription(p, lead, days, None, "manual")
        lead.status = "sold"
        lead.save(update_fields=["status"])
        if api:
            S._tell_approved(p, api, lead, type("X", (), {"tariff": sub.tariff})(), sub)
        audit("online.sub_grant", user=request.user, target=f"{p.name}/{lead.full_name}", detail={"days": days})
        messages.success(request, f"{lead.full_name}: obuna {fmt_dt(sub.ends_at)} gacha. Foydalanuvchiga xabar va guruh havolasi yuborildi.")
    elif a == "cancel":
        S.cancel_subscription(p, api, lead)
        audit("online.sub_cancel", user=request.user, target=f"{p.name}/{lead.full_name}")
        messages.success(request, f"{lead.full_name}: obuna bekor qilindi, guruhdan chiqarildi.")
    elif a == "exempt":
        lead.exempt = not lead.exempt
        lead.save(update_fields=["exempt"])
        messages.success(request, f"{lead.full_name}: «cheksiz/xodim» belgisi {'qo‘yildi' if lead.exempt else 'olib tashlandi'}.")
    elif a == "resend" and api:
        if S.has_access(p, lead, lead.tg_id):
            links = S.invite_links(p, api, lead)
            S._send_user(api, lead, D.get_flow(p)["texts"]["join_links"], {"inline_keyboard": [[{"text": f"➡️ {t}", "url": u}] for t, u in links]} if links else None)
            messages.success(request, "Guruh havolalari qayta yuborildi.")
        else:
            messages.error(request, "Faol obuna yo'q.")
    elif a == "kick" and api:
        n = S.remove_from_groups(p, api, lead)
        messages.success(request, f"{n} ta guruhdan chiqarildi.")
    return _back(request, p)


# ───────────────────────── to'lovlar ─────────────────────────
@login_required
@require_permission("online.view")
def online_payments(request, pk):
    p = _proj(pk)
    st = request.GET.get("st", "")
    qs = p.payments.select_related("lead", "tariff")
    if st in ("pending", "approved", "rejected"):
        qs = qs.filter(status=st)
    if request.GET.get("export") == "1" and _can(request, "online.export"):
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "Tolovlar"
        ws.append(["#", "Sana", "Lead", "Telefon", "Tarif", "Summa", "Holat", "Tasdiqlagan", "Sabab"])
        for x in qs:
            ws.append([x.pk, fmt_dt(x.created_at), x.lead.full_name, x.lead.phone, x.tariff.name if x.tariff else "", x.amount, x.get_status_display(), x.decided_by, x.reject_reason])
        buf = io.BytesIO()
        wb.save(buf)
        r = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        r["Content-Disposition"] = f'attachment; filename="tolovlar_{p.pk}.xlsx"'
        return r
    page = Paginator(qs, 40).get_page(request.GET.get("page"))
    return render(request, "online_payments.html", _ctx(request, p, "online_payments", page=page, st=st, can_pay=_can(request, "online.payments"), can_export=_can(request, "online.export"),
                                                        s=S.stats(p)))


@login_required
@require_POST
@require_permission("online.payments")
def online_payment_decide(request, pk, payid):
    p = _proj(pk)
    ok, msg = S.decide_payment(p, _api(p), payid, request.POST.get("do") == "ok", _by(request), request.POST.get("reason", "").strip())
    (messages.success if ok else messages.error)(request, msg)
    audit("online.payment_decided", user=request.user, target=f"{p.name}#{payid}", detail={"do": request.POST.get("do")})
    return _back(request, p, "online_payments")


@login_required
@require_permission("online.view")
def online_receipt(request, pk, payid):
    p = _proj(pk)
    x = get_object_or_404(Payment, pk=payid, project=p)
    api = _api(p)
    if not (api and x.receipt_file_id):
        raise Http404
    try:
        data = api.file_bytes(x.receipt_file_id)
    except Exception:  # noqa: BLE001
        raise Http404
    return HttpResponse(data, content_type="image/jpeg")


# ───────────────────────── tariflar ─────────────────────────
@login_required
@require_permission("online.view")
def online_tariffs(request, pk):
    p = _proj(pk)
    if request.method == "POST":
        if not _can(request, "online.edit"):
            return redirect("online_tariffs", p.pk)
        P = request.POST
        a = P.get("action")

        def num(k, d=0):
            try:
                return max(0, int(re.sub(r"\D", "", P.get(k, "")) or d))
            except ValueError:
                return d
        if a == "add" and P.get("name", "").strip():
            Tariff.objects.create(project=p, name=P["name"].strip()[:80], days=max(1, num("days", 30)), price=num("price"), description=P.get("description", "")[:300], sort=num("sort"))
        elif a == "save":
            t = get_object_or_404(Tariff, pk=P.get("id"), project=p)
            t.name, t.days, t.price = P.get("name", t.name).strip()[:80] or t.name, max(1, num("days", t.days)), num("price")
            t.description, t.sort, t.is_active = P.get("description", "")[:300], num("sort"), bool(P.get("is_active"))
            t.save()
        elif a == "delete":
            t = get_object_or_404(Tariff, pk=P.get("id"), project=p)
            if Payment.objects.filter(tariff=t).exists() or Subscription.objects.filter(tariff=t).exists():
                t.is_active = False
                t.save(update_fields=["is_active"])
                messages.info(request, "Tarif to'lovlarda ishlatilgan, shuning uchun o'chirilmadi — faolsiz qilindi.")
            else:
                t.delete()
        return redirect("online_tariffs", p.pk)
    return render(request, "online_tariffs.html", _ctx(request, p, "online_tariffs", tariffs=p.tariffs.all(), can_edit=_can(request, "online.edit")))


# ───────────────────────── AI mentor ─────────────────────────
@login_required
@require_permission("online.view")
def online_ai(request, pk):
    p = _proj(pk)
    if request.method == "POST":
        if not _can(request, "online.edit"):
            return redirect("online_ai", p.pk)
        P = request.POST
        a = P.get("action")
        if a == "settings":
            p.ai_enabled = bool(P.get("ai_enabled"))
            p.ai_instructions = P.get("ai_instructions", "")[:3000]
            for k in ("ai_free_daily", "ai_sub_daily"):
                try:
                    setattr(p, k, max(0, min(1000, int(P.get(k) or 0))))
                except ValueError:
                    pass
            p.save()
            messages.success(request, "AI sozlamalari saqlandi.")
        elif a == "doc_add":
            content = P.get("content", "").strip()
            f = request.FILES.get("file")
            if f:
                if f.size > 2 * 1024 * 1024 or not f.name.lower().endswith((".txt", ".md", ".csv")):
                    messages.error(request, "Faqat .txt, .md yoki .csv fayl (2 MB gacha) yuklang. PDF bo'lsa, matnini nusxalab qo'ying.")
                    return redirect("online_ai", p.pk)
                content = f.read().decode("utf-8", "ignore").strip() or content
            if P.get("title", "").strip() and content:
                KnowledgeDoc.objects.create(project=p, title=P["title"].strip()[:160], content=content[:60000])
                messages.success(request, "Bilimlar bazasiga qo'shildi.")
            else:
                messages.error(request, "Sarlavha va matn kerak.")
        elif a == "doc_save":
            d = get_object_or_404(KnowledgeDoc, pk=P.get("id"), project=p)
            d.title, d.content, d.is_active = P.get("title", d.title)[:160], P.get("content", d.content)[:60000], bool(P.get("is_active"))
            d.save()
        elif a == "doc_delete":
            KnowledgeDoc.objects.filter(pk=P.get("id"), project=p).delete()
        return redirect("online_ai", p.pk)
    from apps.online.models import AiLog
    return render(request, "online_ai.html", _ctx(request, p, "online_ai", docs=p.docs.all(), can_edit=_can(request, "online.edit"),
                                                  logs=AiLog.objects.filter(project=p, test=False).select_related("lead")[:15], n_today=AiLog.objects.filter(project=p, test=False, at__date=timezone.localdate()).count()))


@login_required
@require_POST
@require_permission("online.edit")
def online_ai_test(request, pk):
    p = _proj(pk)
    try:
        q = (json.loads(request.body.decode() or "{}").get("q") or "").strip()
    except json.JSONDecodeError:
        q = ""
    if not q:
        return JsonResponse({"ok": False, "error": "Savol yozing."}, status=400)
    try:
        return JsonResponse({"ok": True, "answer": mentor.mentor_answer(p, None, q, user=request.user, test=True), "kb": mentor.retrieve(p, q)[:600]})
    except Exception as e:  # noqa: BLE001
        return JsonResponse({"ok": False, "error": getattr(e, "user_message_uz", None) or str(e)}, status=200)


# ───────────────────────── guruhlar va adminlar ─────────────────────────
def _resolve_chat(p, raw):
    raw = (raw or "").strip()
    if not raw:
        return None, ""
    api = _api(p)
    if re.fullmatch(r"-?\d{5,}", raw):
        cid = int(raw)
        title = ""
        if api:
            try:
                title = api.get_chat(cid).get("title", "")
            except TgError:
                pass
        return cid, title
    if api:
        ch = api.get_chat(raw if raw.startswith("@") else "@" + raw.lstrip("@"))
        return int(ch["id"]), ch.get("title", "")
    raise TgError("Bot tokeni kiritilmagan")


@login_required
@require_permission("online.view")
def online_groups(request, pk):
    p = _proj(pk)
    if request.method == "POST":
        if not _can(request, "online.configure"):
            return redirect("online_groups", p.pk)
        P = request.POST
        a = P.get("action")
        api = _api(p)
        try:
            if a == "group_add":
                cid, title = _resolve_chat(p, P.get("chat"))
                if cid:
                    g, _ = ProjectGroup.objects.get_or_create(project=p, chat_id=cid, defaults={"title": title})
                    if api:
                        S.check_group(p, api, g)
            elif a == "group_check" and api:
                S.check_group(p, api, get_object_or_404(ProjectGroup, pk=P.get("id"), project=p))
            elif a == "group_del":
                ProjectGroup.objects.filter(pk=P.get("id"), project=p).delete()
            elif a in ("leads_chat", "pay_chat"):
                cid, title = _resolve_chat(p, P.get("chat"))
                if a == "leads_chat":
                    p.leads_chat_id, p.leads_chat_title = cid, title
                else:
                    p.pay_chat_id, p.pay_chat_title = cid, title
                p.save()
                if cid and api:
                    ok, msg, t2 = S.check_post_chat(api, cid, p.bot_id or api.get_me()["id"])
                    (messages.success if ok else messages.warning)(request, f"{title or cid}: {msg}")
            elif a == "admin_add":
                tg = int(re.sub(r"\D", "", P.get("tg_id", "")) or 0)
                if tg:
                    ProjectAdmin.objects.update_or_create(project=p, tg_id=tg, defaults={"name": P.get("name", "").strip()[:80], "is_primary": bool(P.get("is_primary"))})
                else:
                    messages.error(request, "Telegram ID raqamini kiriting (adminning o'zi botga /start bosib, @userinfobot dan bilib oladi).")
            elif a == "admin_toggle":
                ad = get_object_or_404(ProjectAdmin, pk=P.get("id"), project=p)
                ad.is_primary = not ad.is_primary
                ad.save(update_fields=["is_primary"])
            elif a == "admin_del":
                ProjectAdmin.objects.filter(pk=P.get("id"), project=p).delete()
            audit("online.groups_changed", user=request.user, target=p.name, detail={"action": a})
        except TgError as e:
            messages.error(request, f"Telegram xatosi: {e.description}")
        return redirect("online_groups", p.pk)
    return render(request, "online_groups.html", _ctx(request, p, "online_groups", groups=p.groups.all(), admins=p.admins.all(), can_cfg=_can(request, "online.configure")))


# ───────────────────────── tarix va sozlamalar ─────────────────────────
@login_required
@require_permission("online.view")
def online_events(request, pk):
    p = _proj(pk)
    kind = request.GET.get("k", "")
    qs = p.events.select_related("lead")
    if kind:
        qs = qs.filter(kind=kind)
    kinds = list(p.events.values_list("kind", flat=True).distinct())
    return render(request, "online_events.html", _ctx(request, p, "online_events", page=Paginator(qs, 60).get_page(request.GET.get("page")), kinds=kinds, k=kind))


@login_required
@require_permission("online.configure")
def online_settings(request, pk):
    p = _proj(pk)
    if request.method == "POST":
        P = request.POST
        if P.get("action") == "delete":
            if P.get("confirm", "").strip() == p.name:
                audit("online.project_deleted", user=request.user, target=p.name)
                p.delete()
                messages.success(request, "Loyiha o'chirildi.")
                return redirect("online_home")
            messages.error(request, "O'chirish uchun loyiha nomini aynan yozing.")
            return redirect("online_settings", p.pk)
        p.name = P.get("name", p.name).strip()[:120] or p.name
        if P.get("token", "").strip():
            p.set_token(P["token"])
            try:
                me = _api(p).get_me()
                p.bot_id, p.bot_username = me["id"], me.get("username", "")
            except TgError as e:
                messages.warning(request, f"Token saqlandi, lekin tekshirib bo'lmadi: {e.description}")
        p.card_number = re.sub(r"[^\d ]", "", P.get("card_number", ""))[:32]
        p.card_holder, p.card_bank = P.get("card_holder", "").strip()[:80], P.get("card_bank", "").strip()[:60]
        for k, lo, hi in (("escalate_min", 1, 1440), ("repost_days", 1, 365)):
            try:
                setattr(p, k, max(lo, min(hi, int(P.get(k) or 0))))
            except ValueError:
                pass
        p.remind_days = ",".join(x for x in re.split(r"[,\s]+", P.get("remind_days", "")) if x.isdigit())[:20] or "3,1"
        p.save()
        audit("online.settings_changed", user=request.user, target=p.name)
        messages.success(request, "Sozlamalar saqlandi.")
        return redirect("online_settings", p.pk)
    return render(request, "online_settings.html", _ctx(request, p, "online_settings", has_token=bool(p.token_enc)))
