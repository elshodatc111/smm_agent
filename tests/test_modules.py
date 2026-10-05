import json
import subprocess
from datetime import date

import pytest
from django.utils import timezone

from apps.metrics import collector


def test_studio_flows_with_fake_ai(biz, acc_tg, admin):
    from apps.studio import services as st
    collector.collect_account(acc_tg)
    # kontent-reja: fake AI bo'sh yozuv beradi -> xatosiz ishlashi kerak
    plan = st.monthly_plan(biz, "telegram", date.today(), admin)
    assert plan.pk and plan.entries == []
    plan.entries = [{"date": date.today().isoformat(), "format": "reel", "topic": "Yangi kurs", "hook": "3 soniyada...", "cta": "Yoziling", "goal": "lead"}]
    plan.save()
    assert st.load_plan_to_calendar(plan, admin) == 1
    from apps.content.models import ContentItem
    it = ContentItem.objects.get(plan=plan)
    assert it.status == "draft" and it.ai_draft and it.targets.count() == 1
    s = st.generate_script(biz, "kurs reklamasi", "reel", user=admin)
    assert s.pk
    a = st.analyze_script(s, admin)
    assert a.pk
    p = biz.posts.first()
    ins = st.why_worked(p, admin)
    assert ins.numbers and ins.verdict in ("winner", "average", "weak")
    from apps.metrics.models import Comment
    c = Comment.objects.first()
    st.classify_comment(c)
    sug = st.suggest_reply(c, admin)
    assert sug.draft
    assert st.refresh_hook_suggestions(biz) >= 1


def test_brand_profile_versioning(biz, admin):
    from apps.studio import services as st
    from apps.businesses.services import approve_brand, restore_brand
    bp = st.generate_brand_profile(biz, admin)
    assert bp.version == 1 and bp.status == "draft" and bp.generated_by_ai
    assert biz.active_brand is None
    approve_brand(bp, admin)
    assert biz.active_brand.version == 1
    from apps.businesses.services import save_brand_version
    save_brand_version(biz, "yangi matn", user=admin, status="approved")
    assert biz.active_brand.version == 2
    new = restore_brand(bp, admin)
    assert new.version == 3 and new.content == bp.content


def test_report_generation_pdf_xlsx(biz, acc_tg, admin):
    from apps.reports.generator import default_period, generate_report
    from apps.reports.models import Report
    collector.collect_account(acc_tg)
    ps, pe = default_period("monthly")
    r = Report.objects.create(business=biz, kind="monthly", period_start=ps, period_end=pe, created_by=admin)
    generate_report(r)
    r.refresh_from_db()
    assert r.status == "ready"
    with open(r.pdf.path, "rb") as f:
        assert f.read(5) == b"%PDF-"
    from openpyxl import load_workbook
    wb = load_workbook(r.xlsx.path)
    assert "Umumiy" in wb.sheetnames and "Tavsiyalar" in wb.sheetnames
    assert r.summary_text


def test_report_delivery_via_bot(biz, acc_tg, admin):
    from apps.reports.generator import default_period, generate_report
    from apps.reports.models import Report
    from apps.tgbot import bot
    from apps.tgbot.models import ClientBotUser
    ps, pe = default_period("weekly")
    r = Report.objects.create(business=biz, kind="weekly", period_start=ps, period_end=pe)
    generate_report(r)
    ClientBotUser.objects.create(tg_user_id=1, business_ids=[biz.pk], current_business=biz.pk)
    ClientBotUser.objects.create(tg_user_id=2, business_ids=[], current_business=None)
    api = bot.FakeBotAPI()
    assert bot.deliver_report(r, api) == 1 and api.docs[0][0] == 1


def test_links_roundtrip_with_fake_worker(biz, admin):
    from apps.links import services as ls
    from apps.links.models import LinkClick
    link = ls.create_link(biz, "https://example.com/kurs", utm={"source": "telegram", "campaign": "oktabr"})
    assert link.synced and ls.FakeWorker.links[link.slug]["url"].startswith("https://example.com/kurs?")
    assert "utm_source=telegram" in ls.FakeWorker.links[link.slug]["url"]
    ls.FakeWorker.clicks.append({"id": 1, "slug": link.slug, "t": int(timezone.now().timestamp()), "source": "t.me", "device": "mobile", "country": "UZ"})
    ls.sync_clicks()
    ls.sync_clicks()  # dublikat yo'q
    assert LinkClick.objects.count() == 1
    ls.remove_link(link)
    assert link.slug not in ls.FakeWorker.links


def test_health_snapshot(biz, acc_tg):
    from apps.health import services
    snap = services.snapshot()
    assert {"server", "queue", "providers", "ai", "connections", "r2", "worker", "tips"} <= set(snap)
    assert any("Worker" in t for t in snap["tips"])


def test_backup_sqlite(settings, tmp_path, biz, monkeypatch):
    from apps.health import jobs
    settings.BASE_DIR = tmp_path
    settings.DATABASES["default"]["ENGINE"] = "django.db.backends.sqlite3"
    dbf = tmp_path / "x.sqlite3"
    dbf.write_bytes(b"data")
    settings.DATABASES["default"]["NAME"] = str(dbf)
    rec = jobs.run_backup()
    assert rec.ok and rec.size > 0


def test_media_pipeline_on_synthetic_video(biz, admin, tmp_path):
    from django.core.files.base import ContentFile
    from apps.media_analysis.models import MediaFile
    from apps.media_analysis.pipeline import analyze_media, sha256_file
    vid = tmp_path / "t.mp4"
    r = subprocess.run(["ffmpeg", "-y", "-f", "lavfi", "-i", "testsrc=duration=4:size=360x640:rate=15", "-f", "lavfi", "-i", "sine=frequency=440:duration=4",
                        "-shortest", "-pix_fmt", "yuv420p", str(vid)], capture_output=True)
    assert r.returncode == 0
    m = MediaFile(business=biz, kind="video", original_name="t.mp4", uploaded_by=admin, size=vid.stat().st_size, sha256=sha256_file(str(vid)))
    m.file.save("t.mp4", ContentFile(vid.read_bytes()))
    m.save()
    a = analyze_media(m)
    assert a.technical["duration"] > 3 and a.technical["aspect"] < 0.6
    assert a.transcript and a.transcript["text"]
    assert a.ai is not None and not a.degraded
    # xesh bo'yicha kesh
    m2 = MediaFile(business=biz, kind="video", sha256=m.sha256)
    m2.file.save("t2.mp4", ContentFile(vid.read_bytes()))
    m2.save()
    b = analyze_media(m2)
    assert b.step_results.get("cached_from") == a.pk


def test_stt_confidence_and_uz_params():
    from apps.media_analysis.pipeline import confidence_from
    assert confidence_from({"avg_logprob": -0.2, "no_speech_prob": 0.0}) > 0.75
    assert confidence_from({"avg_logprob": -1.4, "no_speech_prob": 0.4}) < 0.2


def test_instagram_connector_with_mocked_http(acc_ig, monkeypatch, settings):
    settings.USE_FAKES = False
    from apps.connections import instagram
    from apps.connections.base import MediaType, PublishRequest
    acc_ig.set_credentials({"access_token": "TOKEN123456789", "ig_user_id": "17841"})
    acc_ig.save()
    calls = []

    class R:
        def __init__(self, j, status=200): self._j, self.status_code, self.headers, self.text, self.content = j, status, {}, json.dumps(j), b"1"
        def json(self): return self._j

    def fake_request(method, url, *, key, **kw):
        calls.append((method, url, kw.get("data") or kw.get("params")))
        if url.endswith("/17841/media") and method == "POST":
            return R({"id": "C1", "uri": "https://rupload.facebook.com/x"} if (kw["data"].get("upload_type") == "resumable") else {"id": "C1"})
        if "rupload" in url:
            return R({"success": True})
        if url.endswith("/C1"):
            return R({"status_code": "FINISHED"})
        if url.endswith("/media_publish"):
            return R({"id": "M1"})
        if url.endswith("/M1"):
            return R({"permalink": "https://instagram.com/p/x", "timestamp": "2026-10-05T10:00:00+0000"})
        return R({})
    monkeypatch.setattr(instagram, "request", fake_request)
    c = instagram.InstagramConnector(acc_ig)
    # rasm: ochiq URL kerak
    res = c.publish(PublishRequest("k", MediaType.IMAGE, "salom", [], media_public_urls=["https://tmp/x.jpg"]))
    assert res.external_id == "M1"
    assert any("image_url" in str(c_[2]) for c_ in calls)
    # video: resumable (URL shart emas)
    v = tmp = None
    import tempfile, os
    f = tempfile.NamedTemporaryFile(suffix=".mp4", delete=False); f.write(b"0" * 100); f.close()
    calls.clear()
    c.publish(PublishRequest("k2", MediaType.REEL, "reel", [f.name]))
    os.unlink(f.name)
    assert any("rupload" in u for _, u, _ in calls) and not any("video_url" in str(d) for _, _, d in calls)
    # rasm uchun URL bo'lmasa tushunarli xato
    from apps.connections.base import ContentRejectedError
    with pytest.raises(ContentRejectedError):
        c.publish(PublishRequest("k3", MediaType.IMAGE, "x", ["/tmp/x.jpg"]))


def test_public_tme_parser():
    from apps.connections.public_tme import parse_channel_html
    html = '''<div class="tgme_channel_info_counter"><span class="counter_value">12.5K</span> <span class="counter_type">subscribers</span></div>
    <div class="tgme_widget_message_wrap js-widget_message_wrap"><div class="tgme_widget_message" data-post="kanal/15">
    <div class="tgme_widget_message_text js-message_text">Salom <b>dunyo</b></div><span class="tgme_widget_message_views">1.2K</span>
    <time datetime="2026-10-01T10:00:00+00:00"></time></div></div>'''
    i = parse_channel_html(html)
    assert i["followers"] == 12500 and i["posts"][0]["views"] == 1200 and i["posts"][0]["id"] == "15"
